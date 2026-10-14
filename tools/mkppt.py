"""Simple white SIH-style Datathon Round 1 PPT template. Usage: python mkppt.py <assets_dir> <out.pptx>"""
import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

A, OUT = sys.argv[1], sys.argv[2]
BLACK = RGBColor(0, 0, 0); WHITE = RGBColor(255, 255, 255)
CYAN = RGBColor(0x19, 0xE3, 0xF5); WASH = RGBColor(0xDF, 0xFB, 0xFE); GREY = RGBColor(0x55, 0x55, 0x55)
DISPLAY, BODY = "Arial Black", "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def bgwhite(s):
    f = s.background.fill; f.solid(); f.fore_color.rgb = WHITE

def rect(s, x, y, w, h, fill=None, line=None, lw=1, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = radius
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh

def txt(s, x, y, w, h, runs, size=20, color=BLACK, font=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None):
    tb = shape if shape is not None else s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    if isinstance(runs, str): runs = [(runs, {})]
    for t, o in runs:
        r = p.add_run(); r.text = t; f = r.font
        f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold); f.name = o.get('font', font)
        f.color.rgb = o.get('color', color)
        if o.get('u'): f.underline = True
    return tb

def bullets(s, x, y, w, items, size=24, gap=10):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.5 * len(items) + 0.5))
    tf = tb.text_frame; tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(int(Inches(0.4)))); pPr.set('indent', str(-int(Inches(0.4))))
        bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', 'Arial')
        bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', '•')
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.name = BODY; r.font.color.rgb = BLACK
    return tb

def heading(s, y, text, mandatory=False):
    runs = [("❖ ", {'color': BLACK}), (text, {'u': True})]
    if mandatory: runs.append(("   (Mandatory)", {'size': 16, 'color': GREY, 'bold': False}))
    txt(s, 0.6, y, 12.1, 0.6, runs, size=28, bold=True, color=BLACK)

def chrome(s, n, title):
    bgwhite(s)
    txt(s, 1.95, 0.25, 9.4, 0.95, title, size=30, font=DISPLAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    s.shapes.add_picture(os.path.join(A, "logo-dsc-light.png"), Inches(0.45), Inches(0.15), height=Inches(1.0))
    s.shapes.add_picture(os.path.join(A, "logo-iiec-light.png"), Inches(11.2), Inches(0.35), height=Inches(0.62))
    rect(s, 0, 6.95, 13.333, 0.55, fill=CYAN)
    txt(s, 0, 6.95, 13.333, 0.55, "DATATHON 2026  ·  Round 1 Idea Submission – Template", size=12, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 11.9, 6.95, 1.1, 0.55, str(n), size=13, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

# 1 ---------- title
s = prs.slides.add_slide(blank); bgwhite(s)
s.shapes.add_picture(os.path.join(A, "logo-dsc-light.png"), Inches(0.55), Inches(0.25), height=Inches(1.15))
s.shapes.add_picture(os.path.join(A, "logo-vitap-light.png"), Inches(4.9), Inches(0.4), height=Inches(0.85))
s.shapes.add_picture(os.path.join(A, "logo-iiec-light.png"), Inches(10.75), Inches(0.5), height=Inches(0.75))
rect(s, 0.5, 1.55, 12.33, 0.04, fill=CYAN)
txt(s, 0.5, 1.7, 12.3, 1.5, [("DATA", {}), ("T", {'color': CYAN}), ("HON 2026", {})], size=66, font=DISPLAY, align=PP_ALIGN.CENTER)
txt(s, 0.5, 3.05, 12.3, 0.6, "ROUND 1  ·  IDEA PPT + BASE PROTOTYPE", size=24, bold=True, align=PP_ALIGN.CENTER)
bullets(s, 3.2, 3.9, 8.5, ["Team Name (as registered)", "Team Members (1–2)", "Track Chosen", "Problem Statement Title", "Product Name"], size=24, gap=8)

# content slides ---------- (heading + short bullets, SIH style)
def slide(n, title, blocks):
    s = prs.slides.add_slide(blank); chrome(s, n, title)
    y = 1.65
    for text, items, mand in blocks:
        heading(s, y, text, mand)
        bullets(s, 0.6, y + 0.75, 12.1, items)
        y += 0.75 + 0.62 * len(items) + 0.55
    return s

slide(2, "PROBLEM & TARGET USER", [
    ("Problem & Target User", ["Who is the user and what is their pain point", "Why this problem matters", "Evidence from the dataset"], False)])
slide(3, "SOLUTION & USP", [
    ("Proposed Solution", ["Describe your idea / solution / prototype", "How it addresses the problem"], False),
    ("USP", ["What makes your solution unique", "Innovation compared to the obvious approach"], True)])
slide(4, "DATASET & INSIGHTS", [
    ("Dataset Usage", ["Tables used and how they are joined", "Cleaning done and assumptions made", "Key insights found"], False)])
slide(5, "ML MODEL", [
    ("ML Approach", ["Model built from scratch, or using an API key?", "Features, target, baseline, split and metric", "How the model output is used in the product"], False)])
slide(6, "BASE PROTOTYPE & MVP", [
    ("Base Prototype (Round 1)", ["What is built so far (screenshots / link)"], False),
    ("MVP for Round 2", ["Must-have features for the complete prototype", "Demo workflow: input → output → action"], True)])
slide(7, "TECHNICAL APPROACH", [
    ("Technical Approach", ["Technologies used (languages, frameworks, APIs)", "Flow: Data → Analyze → Model → Build → Deploy"], False)])
slide(8, "FEASIBILITY, IMPACT & NEXT STEPS", [
    ("Feasibility, Impact & Next Steps", ["Feasibility, risks and how to overcome them", "Impact and limitations (synthetic data)", "Next steps and references"], False)])

prs.save(OUT)
print("saved", OUT)
