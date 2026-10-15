"""Judges' sheets: one-page landscape Word doc per round + one Excel workbook (Round 1 / Round 2 sheets).
Usage: python make_scoresheets.py <out_dir>"""
import sys, os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)

TRACKS = ["01 Health", "02 Education", "03 Smart City", "04 Business"]
ROUNDS = {
    1: dict(title="Round 1 · PPT + Base Prototype", file="DATATHON-Round1-Judge-Sheet.docx", shortlist=False,
            crit=[("Problem & target user", 25), ("Idea / solution & USP", 40), ("Feasibility & MVP plan", 25), ("Base prototype", 10)],
            widths=[1.0, 6.0, 2.8, 2.7, 2.7, 2.7, 2.4, 2.4, 4.0]),
    # Round 2: website weights grouped into 5 segments (10+10, 15+25, 20, 10+5, 5)
    2: dict(title="Round 2 · Complete Prototype", file="DATATHON-Round2-Judge-Sheet.docx", shortlist=False,
            crit=[("Problem & data use", 20), ("ML model (from scratch scores most)", 40), ("Working product", 20), ("USP, impact & feasibility", 15), ("Demo & pitch", 5)],
            widths=[1.0, 5.4, 2.5, 2.4, 3.0, 2.4, 2.6, 2.0, 2.2, 3.2]),
}
for r in ROUNDS.values():
    assert sum(m for _, m in r["crit"]) == 100

BLACK, CYAN, WHITE = "000000", "19E3F5", "FFFFFF"

# ====================================================================== WORD
def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_)
    tcPr.append(shd)

def cell_text(cell, text, size=8, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]; p.alignment = align
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.font.name = "Arial"
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = 1

def set_widths(table, widths_cm):
    table.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); table._tbl.tblPr.append(lay)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)

def row_flag(row, tag):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement(tag); e.set(qn('w:val'), 'true'); trPr.append(e)

def make_doc(cfg, rows=10, row_h=1.25):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.5); sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.0)
    doc.styles['Normal'].font.name = 'Arial'
    hdr = sec.header.paragraphs[0]; hdr.text = ""
    r = hdr.add_run("DATATHON 2026 · VIT-AP · Data Science Club × IIEC · Judges' Scoresheet"); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(cfg["title"]); r.font.size = Pt(22); r.bold = True; r.font.name = "Arial"

    t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, lab in zip(t.rows[0].cells, ["Judge name:", "Signature:", "Date: 16 Oct 2026"]):
        cell_text(c, lab, 10, True, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_widths(t, [10.0, 9.0, 7.7])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    crit = cfg["crit"]
    head = ["Sl", "Team name", "Track"] + [l for l, _ in crit] + ["Total (100)"] + (["Shortlist Y / N"] if cfg["shortlist"] else []) + ["Remarks"]
    maxes = ["", "", ""] + [str(m) for _, m in crit] + ["100"] + ([""] if cfg["shortlist"] else []) + [""]
    assert len(head) == len(cfg["widths"]) and abs(sum(cfg["widths"]) - 26.7) < 0.05, (len(head), sum(cfg["widths"]))
    tb = doc.add_table(rows=2, cols=len(head)); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (h, m) in enumerate(zip(head, maxes)):
        cell_text(tb.rows[0].cells[i], h, 10, True, WHITE); shade(tb.rows[0].cells[i], BLACK)
        cell_text(tb.rows[1].cells[i], m, 11, True); shade(tb.rows[1].cells[i], CYAN)
    tb.rows[0].height = Cm(1.3); tb.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for n in range(1, rows + 1):
        row = tb.add_row(); row_flag(row, 'w:cantSplit')
        row.height = Cm(row_h); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cell_text(row.cells[0], str(n), 10)
        for c in row.cells[1:]: cell_text(c, "", 10)
    set_widths(tb, cfg["widths"])
    path = os.path.join(OUT, cfg["file"]); doc.save(path); return path

paths = [make_doc(c) for c in ROUNDS.values()]

# ====================================================================== EXCEL
wb = Workbook()
thin = Side(style='thin', color='999999'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
F_BLACK = PatternFill('solid', start_color=BLACK); F_CYAN = PatternFill('solid', start_color=CYAN); F_WASH = PatternFill('solid', start_color='DFFBFE')
H = Font(bold=True, color=WHITE, name='Arial', size=10); B = Font(bold=True, name='Arial', size=10); N = Font(name='Arial', size=10)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True); LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
ROWS = 100

def make_sheet(ws, cfg):
    crit = cfg["crit"]; nc = len(crit)
    head = ["Sl", "Team name", "Track"] + [l for l, _ in crit] + ["Total (100)"] + (["Shortlist Y / N"] if cfg["shortlist"] else []) + ["Remarks"]
    ncol = len(head); fc = 4; lc = fc + nc - 1; tc = lc + 1
    ws["A1"] = "DATATHON 2026 — " + cfg["title"]; ws["A1"].font = Font(bold=True, size=16, name='Arial')
    ws["A2"] = "Judge name:"; ws["A2"].font = B; ws.merge_cells("C2:E2"); ws["C2"].fill = F_WASH
    ws["G2"] = "Date:"; ws["G2"].font = B; ws["H2"] = "16 Oct 2026"; ws["H2"].font = Font(name='Arial', size=10, color='0000FF')
    hr = 4
    for i, t in enumerate(head):
        c = ws.cell(row=hr, column=1 + i, value=t); c.font = H; c.fill = F_BLACK; c.alignment = CEN; c.border = box
        m = ws.cell(row=hr + 1, column=1 + i); m.fill = F_CYAN; m.border = box; m.alignment = CEN; m.font = B
    ws.cell(row=hr + 1, column=1, value="Max")
    for k, (_, mx) in enumerate(crit):
        ws.cell(row=hr + 1, column=fc + k, value=mx).font = Font(name='Arial', size=10, bold=True, color='0000FF')
    f0 = ws.cell(row=hr + 1, column=fc).coordinate; l0 = ws.cell(row=hr + 1, column=lc).coordinate
    ws.cell(row=hr + 1, column=tc, value=f"=SUM({f0}:{l0})")
    ws.row_dimensions[hr].height = 40; ws.row_dimensions[hr + 1].height = 22
    s = hr + 2; e = s + ROWS - 1
    for n in range(ROWS):
        rr = s + n; ws.row_dimensions[rr].height = 26
        ws.cell(row=rr, column=1, value=n + 1).font = N
        for ci in range(1, ncol + 1):
            c = ws.cell(row=rr, column=ci); c.border = box
            c.alignment = LEFT if ci in (2, ncol) else CEN
            if ci > 1 and ci != tc: c.font = Font(name='Arial', size=11, color='0000FF')
        a = ws.cell(row=rr, column=fc).coordinate; b = ws.cell(row=rr, column=lc).coordinate
        t = ws.cell(row=rr, column=tc, value=f'=IF(COUNT({a}:{b})=0,"",SUM({a}:{b}))'); t.font = Font(name='Arial', size=11, bold=True); t.fill = F_WASH; t.alignment = CEN
    for k, (lab, mx) in enumerate(crit):
        col = ws.cell(row=1, column=fc + k).column_letter
        dv = DataValidation(type="decimal", operator="between", formula1="0", formula2=f"${col}${hr+1}", allow_blank=True, showErrorMessage=True,
                            errorTitle="Out of range", error=f"{lab}: enter a mark from 0 to {mx}.")
        ws.add_data_validation(dv); dv.add(f"{col}{s}:{col}{e}")
    dvt = DataValidation(type="list", formula1='"' + ",".join(TRACKS) + '"', allow_blank=True); ws.add_data_validation(dvt); dvt.add(f"C{s}:C{e}")
    if cfg["shortlist"]:
        sc = ws.cell(row=1, column=ncol - 1).column_letter
        dvs = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); ws.add_data_validation(dvs); dvs.add(f"{sc}{s}:{sc}{e}")
    widths = [5, 34, 16] + [18] * nc + [14] + ([14] if cfg["shortlist"] else []) + [34]
    for i, w in enumerate(widths): ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = w
    ws.freeze_panes = ws.cell(row=s, column=3)
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr+1}"; ws.print_area = f"A1:{ws.cell(row=1, column=ncol).column_letter}{e}"
    ws.page_margins.left = ws.page_margins.right = 0.4

ws1 = wb.active; ws1.title = "Round 1"; make_sheet(ws1, ROUNDS[1])
make_sheet(wb.create_sheet("Round 2"), ROUNDS[2])
xlsx = os.path.join(OUT, "DATATHON-Judge-Sheets.xlsx"); wb.save(xlsx)
for p in paths + [xlsx]: print(p)
